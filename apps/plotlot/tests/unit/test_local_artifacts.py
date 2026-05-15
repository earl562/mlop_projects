"""Tests for local document and spreadsheet artifact generation."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from plotlot.retrieval.local_artifacts import create_document, create_spreadsheet


@pytest.fixture
def artifact_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    from plotlot.retrieval import local_artifacts

    monkeypatch.setattr(local_artifacts.settings, "artifact_storage_dir", str(tmp_path))
    return tmp_path


@pytest.mark.asyncio
async def test_create_spreadsheet_writes_xlsx(artifact_dir: Path):
    result = await create_spreadsheet(
        "Vacant Lots",
        ["Address", "Zoning"],
        [["100 Main St", "R-1"], ["200 Oak Ave", "R-2"]],
    )

    path = artifact_dir / result.filename
    assert path.exists()
    assert result.artifact_url == f"/api/v1/artifacts/{result.filename}"
    assert result.content_type.endswith("spreadsheetml.sheet")

    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet["A1"].value == "Address"
    assert sheet["B2"].value == "R-1"


@pytest.mark.asyncio
async def test_create_document_writes_docx(artifact_dir: Path):
    result = await create_document(
        "Zoning Report",
        "Analysis of the R-1 zoning district.\n\nSetbacks require confirmation.",
    )

    path = artifact_dir / result.filename
    assert path.exists()
    assert result.artifact_url == f"/api/v1/artifacts/{result.filename}"
    assert result.content_type.endswith("wordprocessingml.document")

    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Zoning Report" in text
    assert "Analysis of the R-1 zoning district." in text
